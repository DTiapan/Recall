"""Tests for format-aware chunking adapters and table serialization."""

from datetime import datetime
from pathlib import Path
import pytest
import docx

from recall.adapters import ChunkingAdapterRegistry
from recall.adapters.docx_adapter import DocxChunkingAdapter
from recall.adapters.markdown_adapter import MarkdownChunkingAdapter
from recall.adapters.text_adapter import TextChunkingAdapter
from recall.chunkers.table_formatter import TableFormatter
from recall.core.models import IngestConfig


def test_table_formatter_markdown_output():
    formatter = TableFormatter()
    rows = [
        ["Product", "Q1 Revenue", "Growth"],
        ["Cloud Engine", "$120M", "34%"],
        ["AI Platform", "$85M", "62%"],
    ]
    md = formatter.to_markdown(rows, caption="Quarterly Metrics")
    assert "**Table: Quarterly Metrics**" in md
    assert "| Product      | Q1 Revenue | Growth |" in md
    assert "| Cloud Engine | $120M      | 34%    |" in md


def test_table_formatter_chunking_retains_headers():
    formatter = TableFormatter()
    header = ["Department", "Headcount", "Budget"]
    # Build 40 rows to exceed 100 token threshold
    data = [[f"Dept-{i}", f"{10 + i * 2}", f"${1000 * i}"] for i in range(40)]
    rows = [header] + data

    chunks = formatter.chunk_table(rows, max_tokens=100, caption="Enterprise Departments")
    assert len(chunks) > 1

    for chunk in chunks:
        # Every single chunk MUST contain the table headers
        assert "| Department" in chunk
        assert "| Headcount" in chunk
        assert "| Budget" in chunk
        assert "**Table: Enterprise Departments" in chunk


def test_markdown_chunking_adapter(tmp_path: Path):
    file = tmp_path / "compliance_policy.md"
    content = (
        "---\n"
        "policy: Information Security Policy\n"
        "version: 2.1\n"
        "department: Security & Compliance\n"
        "---\n"
        "# Access Controls\n\n"
        "All engineers must use hardware security keys for production SSH access.\n\n"
        "## Key Rotation Matrix\n\n"
        "| Key Type | Max Validity | Approval Required |\n"
        "|---|---|---|\n"
        "| Production Admin | 90 Days | VP Eng |\n"
        "| Staging Access | 180 Days | Team Lead |\n\n"
        "```bash\n"
        "# Rotate SSH key\n"
        "ssh-keygen -t ed25519-sk\n"
        "```\n"
    )
    file.write_text(content, encoding="utf-8")

    adapter = MarkdownChunkingAdapter()
    # Use chunk_size=35 to split across sections and code blocks
    chunks = adapter.chunk(file, config=IngestConfig(chunk_size=35))

    assert len(chunks) >= 2
    types = {c.metadata.content_type for c in chunks}
    assert "text" in types or "table" in types

    # Check provenance metadata
    for c in chunks:
        assert c.metadata.policy_name == "Information Security Policy"
        assert c.metadata.file_type == "markdown"
        assert "Access Controls" in c.metadata.section_hierarchy


def test_docx_chunking_adapter(tmp_path: Path):
    file = tmp_path / "employee_handbook.docx"
    doc = docx.Document()
    doc.add_heading("Employee Benefits", level=1)
    doc.add_paragraph("Comprehensive healthcare is provided to all full-time employees.")
    doc.add_heading("Dental Coverage", level=2)
    doc.add_paragraph("Annual preventive visits are covered at 100%.")

    # Add a table in docx
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Tier"
    table.rows[0].cells[1].text = "Coverage"
    table.rows[1].cells[0].text = "Tier 1"
    table.rows[1].cells[1].text = "In-Network 90%"
    table.rows[2].cells[0].text = "Tier 2"
    table.rows[2].cells[1].text = "Out-of-Network 70%"

    doc.save(file)

    adapter = DocxChunkingAdapter()
    chunks = adapter.chunk(file)

    assert len(chunks) >= 2
    table_chunks = [c for c in chunks if c.metadata.content_type == "table"]
    text_chunks = [c for c in chunks if c.metadata.content_type == "text"]

    assert len(table_chunks) >= 1
    assert "| Tier   | Coverage           |" in table_chunks[0].text

    # Verify heading hierarchy breadcrumb
    dental_chunk = [c for c in text_chunks if "Annual preventive" in c.text][0]
    assert dental_chunk.metadata.section_hierarchy == ["employee_handbook", "Employee Benefits", "Dental Coverage"]
    assert dental_chunk.metadata.file_type == "docx"


def test_adapter_registry(tmp_path: Path):
    registry = ChunkingAdapterRegistry()

    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("Meeting notes from the sprint retrospective.", encoding="utf-8")

    md_file = tmp_path / "guide.md"
    md_file.write_text("# Setup\nRun setup script.", encoding="utf-8")

    txt_chunks = registry.process(txt_file)
    assert len(txt_chunks) == 1
    assert txt_chunks[0].metadata.file_type == "text"

    md_chunks = registry.process(md_file)
    assert len(md_chunks) == 1
    assert md_chunks[0].metadata.file_type == "markdown"


def test_pdf_chunking_adapter(tmp_path: Path):
    from pypdf import PdfWriter
    from recall.adapters.pdf_adapter import PDFChunkingAdapter

    pdf_file = tmp_path / "sample.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with open(pdf_file, "wb") as f:
        writer.write(f)

    adapter = PDFChunkingAdapter()
    chunks = adapter.chunk(pdf_file)
    # A blank page has no text, so chunks list is empty
    assert isinstance(chunks, list)
